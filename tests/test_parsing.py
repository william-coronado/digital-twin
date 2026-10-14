import os

import docx
import pandas as pd
import pytest

from conftest import SAMPLE_DIR
from ingest import parse_file


def parse_sample(rel_path):
    return parse_file(os.path.join(SAMPLE_DIR, rel_path), rel_path)


def write_pdf(path, pages):
    """Writes a minimal valid PDF with one line of Helvetica text per page."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>",
               f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(len(pages)))}] /Count {len(pages)} >>",
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)


def test_markdown_keeps_text_and_relative_source():
    docs = parse_sample("policies/health_safety.md")
    assert len(docs) == 1
    assert docs[0].metadata == {"source": "policies/health_safety.md"}
    assert "Jane Doe" in docs[0].page_content


def test_html_drops_markup():
    [doc] = parse_sample("company_policy.html")
    assert "3 days per week" in doc.page_content
    assert "<p>" not in doc.page_content


def test_csv_rows_become_labelled_lines():
    [doc] = parse_sample("employees.csv")
    lines = doc.page_content.splitlines()
    assert lines[0] == "Row 1: Name: Alice Smith, Role: Director of Product, Email: alice@twin.local, Location: Melbourne"
    assert len(lines) == 3


def test_pdf_gives_one_document_per_page_with_page_numbers(tmp_path):
    write_pdf(tmp_path / "report.pdf", ["Revenue grew in 2022", "Outlook for 2023"])
    docs = parse_file(str(tmp_path / "report.pdf"), "report.pdf")
    assert [d.metadata for d in docs] == [{"source": "report.pdf", "page": 1}, {"source": "report.pdf", "page": 2}]
    assert "Revenue grew" in docs[0].page_content
    assert "Outlook" in docs[1].page_content


def test_docx_includes_paragraphs_and_tables(tmp_path):
    d = docx.Document()
    d.add_paragraph("Quarterly summary")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Revenue", "$5M"
    d.save(tmp_path / "summary.docx")
    [doc] = parse_file(str(tmp_path / "summary.docx"), "summary.docx")
    assert "Quarterly summary" in doc.page_content
    assert "Revenue | $5M" in doc.page_content


def test_xlsx_labels_every_sheet(tmp_path):
    with pd.ExcelWriter(tmp_path / "budget.xlsx") as writer:
        pd.DataFrame({"Item": ["Batteries"], "Cost": [1200]}).to_excel(writer, sheet_name="Hardware", index=False)
        pd.DataFrame({"Item": ["Travel"], "Cost": [300]}).to_excel(writer, sheet_name="Ops", index=False)
    [doc] = parse_file(str(tmp_path / "budget.xlsx"), "budget.xlsx")
    assert "=== Sheet: Hardware ===\nRow 1: Item: Batteries, Cost: 1200" in doc.page_content
    assert "=== Sheet: Ops ===" in doc.page_content


def test_file_without_text_gives_no_documents(tmp_path):
    (tmp_path / "blank.txt").write_text("   \n\n")
    assert parse_file(str(tmp_path / "blank.txt"), "blank.txt") == []


def test_corrupt_file_raises_so_sync_can_record_the_failure(tmp_path):
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    with pytest.raises(Exception):
        parse_file(str(tmp_path / "broken.pdf"), "broken.pdf")
